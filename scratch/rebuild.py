import openpyxl
from openpyxl.worksheet.table import Table, TableStyleInfo
from datetime import datetime

EXCEL_FILE = "LinkedIn_Job_Tracker.xlsx"
HEADERS = [
    'JobID', 'JobURL', 'CompanyName', 'SearchKeyword', 'Status', 
    'ReferralAsked', 'ShortUrlCreated', 'ReferralPerson', 'Remarks', 'CreatedDateTime'
]

jobs = [
    {"JobID": 1, "CompanyName": "PwC Acceleration Center India", "JobURL": "https://tinyurl.com/2bthzb3q", "Status": "Not Interested"},
    {"JobID": 2, "CompanyName": "BNP Paribas", "JobURL": "https://tinyurl.com/2dq7ldpa", "Status": "Not Interested"},
    {"JobID": 3, "CompanyName": "Mphasis", "JobURL": "https://tinyurl.com/26lenb7y", "Status": "Not Interested"},
    {"JobID": 4, "CompanyName": "PwC Acceleration Center India", "JobURL": "https://tinyurl.com/25wja3gu", "Status": "Not Interested"},
    {"JobID": 5, "CompanyName": "Epsilon", "JobURL": "https://tinyurl.com/2873hxe2", "Status": "Ask for referral"},
    {"JobID": 6, "CompanyName": "Fidelity Investments", "JobURL": "https://tinyurl.com/23pjmkyp", "Status": "Not Interested"},
    {"JobID": 7, "CompanyName": "Skan AI", "JobURL": "https://tinyurl.com/2dlqpxka", "Status": "Not Interested"},
    {"JobID": 8, "CompanyName": "CGI", "JobURL": "https://tinyurl.com/2clc9log", "Status": "Not Interested"},
    {"JobID": 9, "CompanyName": "Societe Generale Global Solution Centre", "JobURL": "https://tinyurl.com/2cgbs637", "Status": "Ask for referral"},
    {"JobID": 10, "CompanyName": "Epsilon", "JobURL": "https://tinyurl.com/2dcameeu", "Status": "Ask for referral"},
    {"JobID": 11, "CompanyName": "Blue Yonder", "JobURL": "https://tinyurl.com/2ybddg7m", "Status": "Ask for referral"},
    {"JobID": 12, "CompanyName": "DigiKey Global Capability Center", "JobURL": "https://tinyurl.com/2dytnfs6", "Status": "Not Interested"},
    {"JobID": 13, "CompanyName": "Fidelity Investments", "JobURL": "https://tinyurl.com/249vl23d", "Status": "Not Interested"},
    {"JobID": 14, "CompanyName": "Valtech", "JobURL": "https://tinyurl.com/2bhuqbb9", "Status": "Not Interested"},
    {"JobID": 15, "CompanyName": "Visa", "JobURL": "https://tinyurl.com/2d2px6x5", "Status": "Not Interested"},
    {"JobID": 16, "CompanyName": "Texas Instruments India", "JobURL": "https://tinyurl.com/29hjz668", "Status": "Ask for referral"},
    {"JobID": 17, "CompanyName": "Honeywell", "JobURL": "https://tinyurl.com/29btjagl", "Status": "Ask for referral"}
]

wb = openpyxl.Workbook()
ws = wb.active
ws.title = "Jobs"
ws.append(HEADERS)

created_time = datetime.now().strftime("%Y-%m-%d %H:%M:%S")

for job in jobs:
    row = [
        job["JobID"],
        job["JobURL"],
        job["CompanyName"],
        "SQL Server DBA",
        job["Status"],
        "No",
        "Yes",
        "",
        "",
        created_time
    ]
    ws.append(row)

# Format as an Excel Table
max_row = max(ws.max_row, 1)
ref_range = f"A1:J{max_row}"
tab = Table(displayName="JobTrackerTable", ref=ref_range)
style = TableStyleInfo(
    name="TableStyleLight9", 
    showFirstColumn=False,
    showLastColumn=False, 
    showRowStripes=True, 
    showColumnStripes=False
)
tab.tableStyleInfo = style
ws.add_table(tab)

wb.save(EXCEL_FILE)
print("Successfully reconstructed LinkedIn_Job_Tracker.xlsx!")
